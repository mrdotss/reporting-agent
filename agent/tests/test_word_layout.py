"""Regression coverage for native Word page geometry and readable table flow."""
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

import definition_factory as df
from test_docx import render, row_block


@pytest.mark.parametrize('page_size', ['A4', 'Letter'])
def test_table_grid_fits_the_page_in_twips(page_size):
    _, outcome = render(
        [df.block('table', 'resource_table', {'columns': [df.CPU_AVG, df.CPU_MAX]})],
        design={'page_size': page_size},
    )
    document = Document(BytesIO(outcome.docx_bytes))
    section = document.sections[0]
    budget = section.page_width.twips - section.left_margin.twips - section.right_margin.twips
    table = document.tables[0]
    widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid]
    assert 0 <= budget - sum(widths) < len(widths)
    assert table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w')) == str(budget)


def test_nested_table_fits_its_layout_cell():
    _, outcome = render([row_block('row', [
        [df.block('left', 'resource_table', {'columns': [df.CPU_AVG]})],
        [df.block('right', 'resource_table', {'columns': [df.CPU_MAX]})],
    ])])
    document = Document(BytesIO(outcome.docx_bytes))
    for cell in document.tables[0].rows[0].cells:
        nested = cell.tables[0]
        assert sum(col.width.twips for col in nested.columns) <= cell.width.twips - 200


@pytest.mark.parametrize('density', ['compact', 'normal', 'relaxed'])
def test_rows_stay_whole_without_chaining_entire_table(density):
    compiled, outcome = render(
        [df.block('table', 'resource_table', {'columns': [df.CPU_AVG]})],
        design={'density': density},
    )
    from reporting_agent.render.themes import THEME_SPECS

    table = Document(BytesIO(outcome.docx_bytes)).tables[0]
    for index, row in enumerate(table.rows):
        assert row._tr.trPr.find(qn('w:cantSplit')) is not None
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                assert paragraph.paragraph_format.keep_with_next == (index == 0)
                for run in paragraph.runs:
                    assert run.font.size == Pt(THEME_SPECS['editorial'].small_pt)
    figure_runs = [run.text for row in table.rows for cell in row.cells
                   for p in cell.paragraphs for run in p.runs if run.style.name == 'Figure']
    assert sorted(figure_runs) == sorted(entry.formatted for entry in compiled.ledger.entries.values())


def test_nested_chart_fits_its_layout_cell():
    _, outcome = render([row_block('row', [
        [df.block('chart', 'timeseries_chart', {'metrics': [df.CPU_AVG]})],
        [df.block('text', 'rich_text', {'text': 'Notes'})],
    ])])
    document = Document(BytesIO(outcome.docx_bytes))
    cell = document.tables[0].rows[0].cells[0]
    extent = cell._tc.xpath('.//wp:extent')
    assert extent, 'fixture must contain a rendered chart'
    assert int(extent[0].get('cx')) <= cell.width - 200 * 635


def test_text_facts_use_body_face_without_losing_verification_style():
    from reporting_agent.compile.figures import FigureLedger
    from reporting_agent.render.themes import THEME_SPECS
    from test_text_fact_render import one_fact_table
    from test_text_fact_render import render as render_fact

    ledger = FigureLedger()
    table, fact = one_fact_table(ledger)
    document = Document(BytesIO(render_fact(table, ledger).docx_bytes))
    runs = [run for row in document.tables[0].rows for cell in row.cells
            for paragraph in cell.paragraphs for run in paragraph.runs
            if run.style.name == 'Figure']
    assert [run.text for run in runs] == [fact.formatted]
    assert runs[0].font.name == THEME_SPECS['editorial'].face.body


def test_wide_layout_stacks_without_changing_verified_values():
    compiled, outcome = render([row_block('row', [
        [df.block('wide', 'resource_table', {
            'columns': [df.CPU_AVG, df.CPU_MAX, df.CPU_P95, df.MEMORY_USED_PCT_AVG],
        })],
        [df.block('notes', 'rich_text', {'text': 'Notes after the data'})],
    ])])
    document = Document(BytesIO(outcome.docx_bytes))
    layout = document.tables[0]
    assert len(layout.columns) == 1
    assert len(layout.rows) == 2
    data = layout.rows[0].cells[0].tables[0]
    assert len(data.columns) == 5
    runs = [run.text for row in data.rows for cell in row.cells
            for p in cell.paragraphs for run in p.runs if run.style.name == 'Figure']
    assert sorted(runs) == sorted(entry.formatted for entry in compiled.ledger.entries.values())
    assert 'Notes after the data' in layout.rows[1].cells[0].text


def test_small_layout_keeps_two_columns():
    _, outcome = render([row_block('row', [
        [df.block('left', 'resource_table', {'columns': [df.CPU_AVG]})],
        [df.block('right', 'rich_text', {'text': 'Notes'})],
    ])])
    layout = Document(BytesIO(outcome.docx_bytes)).tables[0]
    assert len(layout.columns) == 2
    assert len(layout.rows) == 1
    assert layout.autofit is False
