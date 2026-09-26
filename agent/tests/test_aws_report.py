"""A whole AWS report: every AWS section, through the real pipeline, to a passing verification.

The fake AWS clients of `test_aws_collector.py` stand in for the account; everything after
them is production code — collection, facts, the AWS section catalogue, compile, the three
renderers, LibreOffice and every verification gate, replay included. Skipped only where
LibreOffice is not installed, exactly as the Azure end-to-end walk is.
"""

from __future__ import annotations

import asyncio
import copy
import json
import os
from io import BytesIO
from typing import Any

import pytest

os.environ.setdefault("AWS_REGION", "us-east-1")
os.environ.setdefault("RPT_ARTIFACT_BUCKET", "rpt-artifacts-test")
os.environ.setdefault("RPT_PROSE_MODEL_ID", "test.prose-model")

from fakes.object_store import InMemoryObjectStore
from reporting_agent.catalog.loader import load_catalog, load_section_catalogue
from reporting_agent.main import StepTracker
from reporting_agent.redaction import discard_secrets
from reporting_agent.render import pdf as P
from reporting_agent.report_pipeline import ReportOutcome, run_generate_report
from test_aws_collector import ACCOUNT, fake_clients, provider
from test_run_end_to_end_v2 import invoke_payload
from test_run_end_to_end_v3 import v3_definition

pytestmark = pytest.mark.skipif(P.SOFFICE_BINARY is None, reason="LibreOffice is not installed")


def aws_definition() -> dict[str, Any]:
    """Every AWS section, each seeded the way the wizard's `addSection` seeds it."""
    # A deep copy: the v2/v3 fixture shares its `design` with every other test's definition,
    # and the number format below is a v3-only field that would invalidate theirs.
    base = copy.deepcopy(v3_definition())
    base["provider"] = "aws"
    sections = []
    for position, entry in enumerate(load_section_catalogue(provider="aws").entries):
        section: dict[str, Any] = {
            "id": f"sec_{entry.key}",
            "type": entry.key,
            "position": position,
            "selection": {
                "resource_types": list(entry.needs_resource_types),
                "resource_groups": [],
                "tag_filters": [],
                "top_n": None,
                "sort": None,
            },
            "metrics": [],
            "presentation": "chart_and_table",
        }
        standard = dict(entry.presets).get("standard_utilization") or ()
        section["metrics"] = [{"metric": metric, "statistic": stat} for metric, stat in standard]
        if entry.key.startswith("historical_"):
            section["lookback"] = 3
        sections.append(section)
    base["sections"] = sections
    base["design"]["number_format"].update({"trim_trailing_zeros": True, "bytes_as_gib": True})
    return base


@pytest.fixture(scope="module")
def walked() -> tuple[InMemoryObjectStore, list[dict[str, Any]], BaseException | None]:
    store = InMemoryObjectStore()
    payload = invoke_payload(aws_definition())
    payload["period"] = {"start": "2026-09-01", "end": "2026-09-01"}
    payload["scope"] = {"resource_types": [], "resource_groups": [], "tag_filters": {}, "regions": ["us-east-1", "ap-southeast-3"]}
    payload.pop("context", None)
    context = {"actor_id": "actor-1", "run_id": "run-1", "subscription_id": ACCOUNT, "timezone": "UTC",
               "provider": "aws", "fidelity_tier": "baseline", "display_name": "Example AWS"}
    events: list[dict[str, Any]] = []
    raised: BaseException | None = None

    async def go() -> None:
        async for event in run_generate_report(
            payload=payload, context=context, steps=StepTracker(), artifact_bucket="unused",
            outcome=ReportOutcome(), prose=None, provider=provider(store, fake_clients()),
            object_store=store, catalog=load_catalog(),
        ):
            events.append(event)

    try:
        asyncio.run(go())
    except Exception as exc:  # a report with gaps ends `PARTIAL_COVERAGE`, which is a delivery
        raised = exc
    finally:
        discard_secrets()
    return store, events, raised


def test_the_report_verifies(walked: tuple[InMemoryObjectStore, list[dict[str, Any]], BaseException | None]) -> None:
    store, events, raised = walked
    verification = [event for event in events if event["type"] == "verification"]
    assert verification and verification[-1]["status"] == "pass", raised
    records = [json.loads(store.get(key).body) for key in store.keys() if "/verification-" in key]  # type: ignore[union-attr]
    assert records and not [f for f in records[-1].get("findings", []) if f.get("severity") == "blocking"]


def test_the_document_speaks_aws(walked: tuple[InMemoryObjectStore, list[dict[str, Any]], BaseException | None]) -> None:
    from docx import Document

    store, _, _ = walked
    key = next(key for key in store.keys() if key.endswith("/report.docx"))
    text = "\n".join(p.text for p in Document(BytesIO(store.get(key).body)).paragraphs)  # type: ignore[union-attr]
    tables = " ".join(cell.text for table in Document(BytesIO(store.get(key).body)).tables  # type: ignore[union-attr]
                      for row in table.rows for cell in row.cells)
    for title in ("AWS Account Overview", "EC2 Instances", "EBS Volumes", "RDS Databases",
                  "EC2 Instance Utilization", "RDS Database Utilization", "EBS Volume Activity",
                  "VPCs and Subnets", "Elastic IP Addresses", "Security Groups in Use", "Backups",
                  "Rightsizing Recommendations"):
        assert title in text, title
    assert "Account ID" in tables and ACCOUNT in tables
    assert "Subscription ID" not in tables and "Resource groups" not in tables
    assert "t3.micro" in tables and "1 GiB" in tables and "admin-sg, web-sg" in tables
    assert "10.0.0.0/16" in tables and "public https" in tables and "203.0.113.10" in tables
    assert "leftover" not in tables  # an unused security group is not reported
